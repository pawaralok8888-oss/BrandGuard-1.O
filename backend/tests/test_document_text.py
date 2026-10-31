"""Reading uploaded guidelines as PDF and DOCX, and the upload route for both.

The PDFs are built here by hand (a valid file with a text layer, or an empty page for a scan), and
the Word files come from python-docx, so no test data files are needed.
"""
from __future__ import annotations

import io

import pytest
from docx import Document
from pypdf import PdfWriter

from app.pipeline.guideline_rules import extract_rule_candidates
from app.services.document_text import UnsupportedDocument, extract_text
from tests.test_guideline_rules import RouteConn, _client, _inserted_rules


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(lines: list[str] | None) -> bytes:
    """A one-page PDF. With lines, each line is a text line. With None, the page has no text (a scan)."""
    if lines is None:
        stream = b""
    else:
        ops = ["BT", "/F1 12 Tf", "72 720 Td"]
        for line in lines:
            ops.append(f"({_pdf_escape(line)}) Tj")
            ops.append("0 -16 Td")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1", errors="replace")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return bytes(out)


def make_docx() -> bytes:
    document = Document()
    document.add_heading("PROHIBITED WORDS", level=1)
    document.add_paragraph('"miracle" (use: "effective" or "clinically proven")', style="List Bullet")
    document.add_paragraph("Professional but approachable.", style="List Bullet")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "cure"
    table.rows[0].cells[1].text = "treat"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------- PDF ----------

def test_pdf_text_layer_is_read_with_its_lines():
    text = extract_text(".pdf", make_pdf(["PROHIBITED WORDS:", '- "miracle" (use: "effective")']))
    assert "PROHIBITED WORDS:" in text
    assert '"miracle"' in text


def test_scanned_pdf_with_no_text_is_refused_with_a_clear_message():
    with pytest.raises(UnsupportedDocument, match="no text layer"):
        extract_text(".pdf", make_pdf(None))


def test_password_protected_pdf_is_refused():
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt(user_password="secret")
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(UnsupportedDocument, match="Password-protected"):
        extract_text(".pdf", buffer.getvalue())


def test_bytes_named_pdf_but_not_a_pdf_are_refused():
    with pytest.raises(UnsupportedDocument, match="not a PDF"):
        extract_text(".pdf", b"plain text pretending")


def test_corrupt_pdf_gives_a_message_not_a_crash():
    with pytest.raises(UnsupportedDocument):
        extract_text(".pdf", b"%PDF-1.4\n garbage that is not a pdf body")


# ---------- DOCX ----------

def test_docx_headings_and_bullets_come_out_in_the_form_the_rule_extractor_reads():
    text = extract_text(".docx", make_docx())
    assert "PROHIBITED WORDS:" in text          # the heading gets the colon it needs
    assert '- "miracle"' in text                 # list paragraphs become bullets
    assert "- cure | treat" in text              # table rows become bullets too


def test_docx_guideline_becomes_draft_rules_end_to_end():
    result = extract_rule_candidates(extract_text(".docx", make_docx()))
    phrases = {c.pattern: c for c in result.candidates if c.kind == "phrase"}
    assert phrases["miracle"].replacement == "effective"
    # Table rows have no quoted term, so they are reported for a person rather than guessed at.
    assert "cure | treat" in result.unparsed
    assert "Professional but approachable." in result.unparsed


def test_docx_that_is_not_a_zip_is_refused():
    with pytest.raises(UnsupportedDocument, match="not a Word document"):
        extract_text(".docx", b"just text")


def test_unknown_extension_is_refused():
    with pytest.raises(UnsupportedDocument):
        extract_text(".exe", b"x")


# ---------- upload route ----------

def test_pdf_upload_saves_its_rules_as_drafts(monkeypatch):
    conn = RouteConn()
    client = _client(conn, monkeypatch)
    pdf = make_pdf(["PROHIBITED WORDS:", '- "miracle" (use: "effective")'])

    res = client.post(
        "/api/knowledge/guidelines/upload",
        files={"file": ("guideline.pdf", pdf, "application/pdf")},
        data={"category": "Terminology"},
    )

    assert res.status_code == 201
    assert res.json()["ruleCandidates"] == 1
    assert len(_inserted_rules(conn)) == 1


def test_scanned_pdf_upload_gives_422_with_the_reason(monkeypatch):
    client = _client(RouteConn(), monkeypatch)
    res = client.post(
        "/api/knowledge/guidelines/upload",
        files={"file": ("scan.pdf", make_pdf(None), "application/pdf")},
        data={"category": "Terminology"},
    )
    assert res.status_code == 422
    assert "no text layer" in res.json()["detail"]


def test_docx_upload_saves_its_rules_as_drafts(monkeypatch):
    conn = RouteConn()
    client = _client(conn, monkeypatch)
    res = client.post(
        "/api/knowledge/guidelines/upload",
        files={"file": ("guideline.docx", make_docx(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"category": "Terminology"},
    )
    assert res.status_code == 201
    assert res.json()["ruleCandidates"] == 1  # only the quoted "miracle" bullet
